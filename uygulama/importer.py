"""CSV / Excel içe aktarma, CSV dışa aktarma ve email doğrulama."""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from database import display_time

EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}$")

# Zorunlu sütunlar ve dosyada karşılık gelebilecek başlıklar.
REQUIRED_COLUMNS = {
    "Email": ["Email", "E-mail", "E-posta", "Eposta", "Mail", "Mail Adresi", "Email Adresi",
              "E-mail Adresi", "E-posta Adresi"],
    "Firma": ["Firma", "Firma Adı", "Şirket", "Sirket", "Şirket Adı", "Kurum", "Company"],
    "Yetkili": ["Yetkili", "Yetkili Kişi", "Yetkili Adı", "İsim", "Isim", "Ad Soyad", "Adı Soyadı",
                "İletişim Kişisi", "Kişi", "Name"],
}


class ImportFailed(Exception):
    pass


@dataclass
class ImportResult:
    rows: list = field(default_factory=list)
    headers: list = field(default_factory=list)
    missing_columns: list = field(default_factory=list)
    invalid_emails: list = field(default_factory=list)  # [(satır no, email)]
    duplicates_in_file: list = field(default_factory=list)
    skipped_empty: int = 0
    skipped_examples: int = 0


def _norm_header(text):
    text = str(text).replace("İ", "i")
    return "".join(text.split()).casefold().replace("-", "").replace("_", "")


def is_valid_email(email):
    return bool(EMAIL_RE.match(email or ""))


def split_emails(text):
    """'a@x.com, b@y.com; c@z.com' -> (geçerliler, geçersizler)"""
    valid, invalid = [], []
    for part in re.split(r"[,;\s]+", text or ""):
        part = part.strip()
        if not part:
            continue
        if is_valid_email(part):
            if part.lower() not in (v.lower() for v in valid):
                valid.append(part)
        else:
            invalid.append(part)
    return valid, invalid


def _cell_to_str(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M") if (value.hour or value.minute) else value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_csv(path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_excel(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportFailed("Excel desteği için 'openpyxl' paketi kurulu değil.") from exc
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFailed("Excel dosyası açılamadı. Dosya bozuk ya da başka bir programda açık olabilir.") from exc
    try:
        sheet = workbook.worksheets[0]
        return [[_cell_to_str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_contacts(path):
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext in (".csv", ".txt"):
            table = _read_csv(path)
        elif ext in (".xlsx", ".xlsm"):
            table = _read_excel(path)
        elif ext == ".xls":
            raise ImportFailed(
                "Eski Excel formatı (.xls) desteklenmiyor.\n"
                "Dosyayı Excel'de açıp 'Farklı Kaydet → Excel Çalışma Kitabı (.xlsx)' ile kaydedin."
            )
        else:
            raise ImportFailed("Sadece .xlsx ve .csv dosyaları destekleniyor.")
    except PermissionError as exc:
        raise ImportFailed("Dosya okunamadı. Excel'de açıksa kapatıp tekrar deneyin.") from exc

    # Baştaki boş satırları atla, ilk dolu satır başlık satırıdır.
    header_index = 0
    while header_index < len(table) and not any(c.strip() for c in table[header_index]):
        header_index += 1
    if header_index >= len(table):
        raise ImportFailed("Dosya boş görünüyor.")

    raw_headers = [h.strip() for h in table[header_index]]
    columns = []  # [(sütun index, başlık)]
    seen = set()
    for i, header in enumerate(raw_headers):
        if header and _norm_header(header) not in seen:
            columns.append((i, header))
            seen.add(_norm_header(header))

    required_index = {}
    for canonical, aliases in REQUIRED_COLUMNS.items():
        alias_keys = {_norm_header(a) for a in aliases}
        for i, header in columns:
            if _norm_header(header) in alias_keys:
                required_index[canonical] = i
                break

    result = ImportResult(headers=[h for _, h in columns])
    result.missing_columns = [c for c in REQUIRED_COLUMNS if c not in required_index]
    if result.missing_columns:
        return result

    required_cols = set(required_index.values())
    seen_emails = set()
    for offset, raw_row in enumerate(table[header_index + 1:]):
        line_no = header_index + offset + 2  # Excel'deki satır numarası
        cells = list(raw_row) + [""] * max(0, len(raw_headers) - len(raw_row))
        if not any(c.strip() for c in cells):
            result.skipped_empty += 1
            continue
        email = cells[required_index["Email"]].strip()
        if email.lower().startswith("mailto:"):
            email = email[7:].strip()
        if not is_valid_email(email):
            result.invalid_emails.append((line_no, email or "(boş)"))
            continue
        if email.lower().endswith("@example.com"):
            result.skipped_examples += 1  # şablondaki örnek satırlar yanlışlıkla gönderilmesin
            continue
        if email.lower() in seen_emails:
            result.duplicates_in_file.append(email)
            continue
        seen_emails.add(email.lower())
        result.rows.append({
            "email": email,
            "firma": cells[required_index["Firma"]].strip(),
            "yetkili": cells[required_index["Yetkili"]].strip(),
            "extra": {h: cells[i].strip() for i, h in columns if i not in required_cols},
        })
    return result


TEMPLATE_HEADERS = ["Email", "Firma", "Yetkili", "Sektör"]
TEMPLATE_EXAMPLES = [
    ["ornek1@example.com", "Örnek Yatırım A.Ş.", "Ayşe Yılmaz", "Finans"],
    ["ornek2@example.com", "Deneme Teknoloji", "Mehmet Kaya", "Teknoloji"],
]
TEMPLATE_HELP = [
    "FİRMA LİSTESİ ŞABLONU - NASIL DOLDURULUR?",
    "",
    "1) 'Firmalar' sayfasının ilk satırındaki başlıkları silmeyin.",
    "2) Zorunlu sütunlar: Email, Firma, Yetkili. Her satır bir firma.",
    "3) Gri renkli örnek satırları silip kendi firmalarınızı yazın.",
    "4) İstediğiniz kadar ek sütun ekleyebilirsiniz (ör. Sektör, Şehir, Unvan).",
    "   Mail metninde {{Sütun Adı}} yazınca her firma için o sütunun değeriyle dolar.",
    "   Örn: 'Sektör' sütunu varsa metinde {{Sektör}} kullanılabilir.",
    "   İhtiyacınız yoksa 'Sektör' sütununu silebilirsiniz.",
    "5) Metinde kullandığınız bir alan bir firmada boşsa, o firmaya mail GÖNDERİLMEZ",
    "   (bozuk mail gitmesin diye). Uygulamada 'Hata: Eksik alan' olarak görünür.",
    "6) Aynı email adresi iki kez yazılırsa sadece ilki alınır.",
    "7) Kaydedip uygulamada 'Dosyadan içe aktar' ile seçin.",
]


def create_template_xlsx(path):
    """Kullanıcının dolduracağı boş Excel firma listesi şablonunu oluşturur."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Firmalar"
    sheet.append(TEMPLATE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A3C6E")
    for row in TEMPLATE_EXAMPLES:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(italic=True, color="808080")
    for column, width in zip("ABCD", (34, 32, 24, 18)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    help_sheet = workbook.create_sheet("Nasıl Doldurulur")
    for line in TEMPLATE_HELP:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True, size=13)
    help_sheet.column_dimensions["A"].width = 95
    workbook.save(path)


REPLY_LABELS = {"replied": "Cevap geldi", "auto": "Otomatik yanıt", "bounced": "Ulaşmadı", "none": "Cevap yok"}


def reply_label(contact):
    if contact.get("manual_reply"):
        return "Cevap alındı (elle)"
    return REPLY_LABELS.get(contact.get("reply_status"), "")


def export_csv(path, contacts, extra_fields):
    """Türkçe Excel'de düzgün açılsın diye ';' ayraçlı ve UTF-8 BOM'lu yazar."""
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["Email", "Firma", "Yetkili", *extra_fields, "Durum", "Gönderim Tarihi", "Hata Mesajı",
                         "Cevap", "Cevap Tarihi", "Hatırlatma Sayısı", "Şablon"])
        for c in contacts:
            writer.writerow([
                c["email"], c["firma"], c["yetkili"],
                *[c["extra"].get(k, "") for k in extra_fields],
                c["status"], display_time(c["sent_at"]), c["error"] or "",
                reply_label(c), display_time(c.get("reply_at")), c.get("reminder_count") or 0,
                c.get("sent_template") or "",
            ])
